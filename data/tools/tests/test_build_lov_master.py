"""Reglas de mapeo del generador del catalogo maestro de LOVs.

Se prueban las funciones puras contra un `aliases.yml` construido en el propio test,
no contra el real: asi un alias nuevo en produccion no rompe estas pruebas, y lo que
se comprueba es el MECANISMO, que es lo que puede romperse al tocar el script.

La ultima clase es la excepcion: contrasta el `data/lov-master.xlsx` ya generado, que
es el fichero que importa la aplicacion. Se salta sola si no esta.

    python3 -m unittest discover -s data/tools/tests
"""

import importlib.util
import os
import unittest

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MASTER = os.path.join(os.path.dirname(BASE), "lov-master.xlsx")

_spec = importlib.util.spec_from_file_location(
    "build_lov_master", os.path.join(BASE, "build_lov_master.py"))
blm = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(blm)


CFG = {
    "code_canonical": {
        "DisconnectorFunction": {
            "FW25": "FW-25",
            "FW+25": "FW-25",
            "F-25": "FW-25",
            "2-25FW": "2xFW-25",
            "LoadB/NZ": "LoadB/NS",
            "DISC/NZ": "Disc/NS",
            "SECT-1": "SECT-I",
        },
    },
    "track_accepted": {
        "DisconnectorFunction": ["FW", "FW-25", "LoadB/t"],
    },
    "code_reassignment": [
        {"from": "DisconnectorFunction",
         "regex": r"^(DISC SECT-I|LOADB/NS FS-1)$", "to": None},
        {"from": "DisconnectorFunction", "regex": r"^(PHQ-\d+|FP)$", "to": None},
        {"from": "Foundation", "regex": "^ANMC?/", "to": "AnchorageFoundation"},
    ],
    "code_rejections": ["N.D."],
}


class FormaCanonicaDeUnCodigo(unittest.TestCase):
    """Varias grafias del mismo equipo tienen que acabar en una sola fila.

    Si no, los usos se reparten entre variantes y cada una queda por debajo del
    umbral de atencion: 'FW-25' con 99 usos se ve, pero partido en 99/14/11/2 no.
    """

    def canonical(self, code, entity="DisconnectorFunction"):
        return blm.canonical_code(entity, code, CFG)

    def test_la_familia_fw_colapsa_en_dos_codigos(self):
        for grafia in ("FW25", "FW+25", "F-25"):
            self.assertEqual(self.canonical(grafia), "FW-25", grafia)
        self.assertEqual(self.canonical("2-25FW"), "2xFW-25")

    def test_nz_y_ns_son_el_mismo_equipo(self):
        # La leyenda de EP9B llama "N.Z. Disconnector" justo al codigo Disc/NS.
        self.assertEqual(self.canonical("LoadB/NZ"), "LoadB/NS")
        self.assertEqual(self.canonical("DISC/NZ"), "Disc/NS")

    def test_la_busqueda_ignora_mayusculas(self):
        # La clave del catalogo ya es insensible a mayusculas; la tabla de alias
        # tiene que serlo tambien o dependeria de como se teclease la celda.
        self.assertEqual(self.canonical("fw25"), "FW-25")
        self.assertEqual(self.canonical("loadb/nz"), "LoadB/NS")

    def test_un_codigo_que_no_esta_en_la_tabla_se_deja_igual(self):
        self.assertEqual(self.canonical("Disc/IO"), "Disc/IO")

    def test_la_tabla_es_por_entidad(self):
        # 'SECT-1' solo es una errata dentro de DisconnectorFunction.
        self.assertEqual(self.canonical("SECT-1"), "SECT-I")
        self.assertEqual(self.canonical("SECT-1", entity="Foundation"), "SECT-1")

    def test_una_entidad_sin_tabla_no_falla(self):
        self.assertEqual(self.canonical("LO QUE SEA", entity="PoleType"), "LO QUE SEA")


class AceptacionDeCodigosDeHojaTrack(unittest.TestCase):
    """La decision de aceptar un codigo de trazado vive en aliases.yml.

    Aceptarlo editando el Excel generado no vale: el maestro se regenera y se lleva
    por delante cualquier ENABLED=SI puesto a mano.
    """

    def accepted(self, code, entity="DisconnectorFunction"):
        return blm.is_track_accepted(entity, code, CFG)

    def test_un_codigo_aceptado_lo_esta(self):
        self.assertTrue(self.accepted("FW"))
        self.assertTrue(self.accepted("LoadB/t"))

    def test_un_codigo_no_aceptado_no_lo_esta(self):
        self.assertFalse(self.accepted("Disc/IO-pr"))

    def test_ignora_mayusculas(self):
        self.assertTrue(self.accepted("loadb/T"))

    def test_una_entidad_sin_lista_no_acepta_nada(self):
        self.assertFalse(self.accepted("FW", entity="PoleType"))


class DescarteDeCeldasQueNoSonUnCodigo(unittest.TestCase):
    """'Sectioning Feeding' es una clave ajena y no puede con dos valores."""

    def route(self, code, entity="DisconnectorFunction"):
        return blm.route_code(entity, code, CFG)

    def test_una_celda_con_dos_equipos_se_descarta_con_motivo(self):
        entity, reason = self.route("Disc SECT-I")
        self.assertIsNone(entity)
        self.assertIn("fuera de catalogo", reason)

    def test_una_fuga_de_otra_columna_se_descarta(self):
        # PHQ-1150 es un brazo y FP un anclaje: se han tecleado en la columna
        # equivocada. Descartarlos con motivo, no cargarlos como alimentacion.
        self.assertIsNone(self.route("PHQ-1150")[0])
        self.assertIsNone(self.route("FP")[0])

    def test_un_codigo_normal_pasa(self):
        self.assertEqual(self.route("Disc/IO"), ("DisconnectorFunction", None))

    def test_las_reglas_son_por_entidad(self):
        # La misma cadena en otra entidad no se toca: la regla dice 'from'.
        self.assertEqual(self.route("FP", entity="Anchorage")[0], "Anchorage")

    def test_un_marcador_sin_valor_se_descarta(self):
        entity, reason = self.route("N.D.")
        self.assertIsNone(entity)
        self.assertIn("marcador", reason)


class DecisionDeEnabledYRevisar(unittest.TestCase):
    """Dos motivos independientes piden revision humana; se comprueban por separado."""

    @staticmethod
    def row(code, sources, entity="DisconnectorFunction"):
        return {"entity": entity, "code": code, "sources": set(sources), "type": ""}

    def test_un_codigo_de_catalogo_curado_entra_sin_revision(self):
        row = self.row("Disc/IO", {"LEGEND", "TRACK"})
        self.assertTrue(blm.decide_enabled(row, CFG))
        self.assertTrue(row["enabled"])
        self.assertFalse(row["revisar"])

    def test_un_codigo_solo_de_track_sin_aceptar_pide_revision(self):
        row = self.row("Disc/IO-pr", {"TRACK"})
        blm.decide_enabled(row, CFG)
        self.assertFalse(row["enabled"])
        self.assertTrue(row["revisar"])

    def test_un_codigo_solo_de_track_ya_aceptado_entra(self):
        row = self.row("FW", {"TRACK"})
        blm.decide_enabled(row, CFG)
        self.assertTrue(row["enabled"])
        self.assertFalse(row["revisar"])

    def test_una_entidad_con_tipo_obligatorio_sin_resolver_pide_revision(self):
        # Foundation exige FoundationType. Sin reglas de tipo no se puede deducir,
        # y cargarlo dejaria una relacion obligatoria a null.
        row = self.row("XXXX", {"BOQ"}, entity="Foundation")
        self.assertFalse(blm.decide_enabled(row, CFG))
        self.assertTrue(row["revisar"])
        self.assertFalse(row["enabled"])


class AgregacionEnElCatalogo(unittest.TestCase):
    """El recorrido completo de una fila: canonicalizar, enrutar y fusionar."""

    def test_dos_grafias_del_mismo_codigo_suman_sus_usos_en_una_fila(self):
        cat = blm.Catalogue(CFG)
        cat.add("DisconnectorFunction", "FW-25", source="TRACK", ep="EP4", track_uses=99)
        cat.add("DisconnectorFunction", "FW25", source="TRACK", ep="EP6", track_uses=14)
        cat.add("DisconnectorFunction", "F-25", source="TRACK", ep="EP6", track_uses=2)

        self.assertEqual(len(cat.rows), 1)
        row = next(iter(cat.rows.values()))
        self.assertEqual(row["code"], "FW-25")
        self.assertEqual(row["track_uses"], 115)
        self.assertEqual(row["eps"], {"EP4", "EP6"})

    def test_lo_descartado_no_entra_en_el_catalogo_pero_queda_registrado(self):
        cat = blm.Catalogue(CFG)
        cat.add("DisconnectorFunction", "Disc SECT-I", source="TRACK", ep="EP7", track_uses=4)

        self.assertEqual(cat.rows, {})
        self.assertEqual(len(cat.discarded), 1)
        self.assertEqual(cat.discarded[0]["codigo"], "Disc SECT-I")

    def test_la_canonicalizacion_va_despues_del_enrutado(self):
        # El alias se busca en la entidad de DESTINO, que es la que decide route_code.
        cat = blm.Catalogue(CFG)
        cat.add("Foundation", "AnM/PL", source="BOQ", ep="EP4")

        row = next(iter(cat.rows.values()))
        self.assertEqual(row["entity"], "AnchorageFoundation")


@unittest.skipUnless(os.path.exists(MASTER), "data/lov-master.xlsx no generado")
class MaestroGenerado(unittest.TestCase):
    """Contraste sobre el fichero real, que es el que importa la aplicacion.

    Cubre lo que las pruebas de mecanismo no ven: que el aliases.yml de produccion
    produce el catalogo que se espera.
    """

    @classmethod
    def setUpClass(cls):
        try:
            import openpyxl
        except ImportError:
            raise unittest.SkipTest("openpyxl no instalado")

        wb = openpyxl.load_workbook(MASTER, read_only=True, data_only=True)
        try:
            sheet = wb["LOVS"]
            header = [str(c) if c else "" for c in
                      next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
            index = {name: n for n, name in enumerate(header)}
            cls.feeding = {
                str(r[1]): r[index["ENABLED"]]
                for r in sheet.iter_rows(min_row=2, values_only=True)
                if r and r[0] == "DisconnectorFunction"
            }
            cls.foundation = {
                str(r[1]): r[index["ENABLED"]]
                for r in sheet.iter_rows(min_row=2, values_only=True)
                if r and r[0] == "Foundation"
            }
            cls.anchorage = {
                str(r[1]): r[index["ENABLED"]]
                for r in sheet.iter_rows(min_row=2, values_only=True)
                if r and r[0] == "Anchorage"
            }
            cls.support_types = {
                str(r[1]): r[index["ENABLED"]]
                for r in sheet.iter_rows(min_row=2, values_only=True)
                if r and r[0] == "SupportType"
            }
            cls.cantilever_types = {
                str(r[1]): (r[index["ENABLED"]], r[index["ORIGEN"]])
                for r in sheet.iter_rows(min_row=2, values_only=True)
                if r and r[0] == "CantileverType"
            }
            cls.steady_arm_types = {
                str(r[1]): r[index["ENABLED"]]
                for r in sheet.iter_rows(min_row=2, values_only=True)
                if r and r[0] == "SteadyArmType"
            }
            cls.portals = {
                str(r[1]): r[index["ENABLED"]]
                for r in sheet.iter_rows(min_row=2, values_only=True)
                if r and r[0] == "Portal"
            }
            cls.foundations = {
                str(r[1]): r[index["ENABLED"]]
                for r in sheet.iter_rows(min_row=2, values_only=True)
                if r and r[0] == "Foundation"
            }
            cls.sectionings = {
                str(r[1]): r[index["ENABLED"]]
                for r in sheet.iter_rows(min_row=2, values_only=True)
                if r and r[0] == "Sectioning"
            }
            cls.foundation_types = {
                str(r[1]): r[3]
                for r in wb["TIPOS"].iter_rows(min_row=2, values_only=True)
                if r and r[0] == "FoundationType"
            }
            cls.unknown_rows = sum(
                1 for r in wb["NO_RECONOCIDO"].iter_rows(min_row=2, values_only=True)
                if r and any(v is not None for v in r))
        finally:
            wb.close()

    def test_los_22_codigos_del_bloque_feeding_estan_habilitados(self):
        # Son los que define la leyenda, identica en los cinco workbooks que la traen.
        legend = {
            "Disc", "Disc/NS", "Disc/IO", "Disc/SI", "Disc/t",
            "LoadB", "LoadB/NS", "LoadB/IO", "LoadB/PP",
            "ED", "ED/T", "SurgeA", "VoltageD", "CurrentT", "SECT-I",
            "FS-1", "FS-1D", "FS/PP-2", "FS/PP-3", "PP-2", "PP-3", "PP-4",
        }
        faltan = {c for c in legend if self.feeding.get(c) != "SI"}
        self.assertEqual(faltan, set())

    def test_la_familia_fw_esta_habilitada_y_canonicalizada(self):
        for code in ("FW", "FW-25", "FW-D", "FW-DV", "2xFW-25"):
            self.assertEqual(self.feeding.get(code), "SI", code)
        for grafia in ("FW25", "FW+25", "F-25", "2-25FW"):
            self.assertNotIn(grafia, self.feeding, grafia)

    def test_las_erratas_de_portico_se_funden_con_el_codigo_bueno(self):
        """'S1PR' y '2PRD' son 'SP1R' y '2PR1D' con las cifras bailadas.

        Se corrigen con un alias y no dandolos de alta: un portico nuevo por cada forma de
        teclear mal el mismo es como el catalogo llego a tener 45 anclajes de los que 17
        no existian.
        """
        self.assertEqual(self.portals.get("SP1R"), "SI")
        self.assertEqual(self.portals.get("2PR1D"), "SI")
        self.assertNotIn("S1PR", self.portals)
        self.assertNotIn("2PRD", self.portals)

    def test_la_cimentacion_p8_es_p8r(self):
        """La P8 solo existe reforzada: 'P8' a secas es la R que falta."""
        self.assertEqual(self.foundations.get("P8R"), "SI")
        self.assertNotIn("P8", self.foundations)

    def test_una_medida_no_entra_como_codigo_de_cimentacion(self):
        """'Ø500*1700' es el diametro y la profundidad de la zapata, no un tipo.

        No se da de alta Y ademas se trata como hueco, para que deje de contarse como una
        decision pendiente: no hay nada que decidir, no es un codigo.
        """
        self.assertNotIn("Ø500*1700", self.foundations)

    def test_el_punto_fijo_doble_y_su_errata_son_el_codigo_MP(self):
        """'2MP' son dos puntos fijos MP y 'MPA' es 'MP' mal tecleado.

        Ninguno de los dos entra como codigo propio. Que los DOS puntos fijos de '2MP' se
        guarden como uno es consecuencia de que el seccionamiento sea un Set, y el Set es
        deliberado: el orden y la repeticion no significan nada en ese campo.
        """
        self.assertEqual(self.sectionings.get("MP"), "SI")
        self.assertNotIn("2MP", self.sectionings)
        self.assertNotIn("MPA", self.sectionings)
        self.assertNotIn("2MP S/A", self.sectionings)
        self.assertNotIn("P50(CS) MPA", self.sectionings)

    def test_las_variantes_nz_se_han_fundido_con_ns(self):
        self.assertNotIn("LoadB/NZ", self.feeding)
        self.assertNotIn("DISC/NZ", self.feeding)

    def test_las_variantes_en_portico_estan_habilitadas(self):
        """'-pr' es el mismo aparato montado en PORTICO, no una errata.

        Lo confirma el propio catalogo: 'LoadB/PP' es «Connections of Load Breaker (without
        feeder)» y 'LoadB/PP-pr' es «... in OCS portal (without feeder)». Lo unico que
        cambia es el portico. Sin habilitarlas, 182 perfiles se quedaban sin su aparato.
        """
        for code in ("Disc-pr", "Disc/IO-pr", "Disc/PP-pr", "Disc/SI-pr",
                     "LoadB-pr", "LoadB/IO-pr", "LoadB/PP-pr"):
            self.assertEqual(self.feeding.get(code), "SI", code)

    def test_el_sufijo_pp_era_una_errata_de_pr(self):
        # Solo aparecia en 4 filas, las cuatro de EP9A, y en ningun BOQ.
        for errata in ("Disc-pp", "LoadB-pp"):
            self.assertNotIn(errata, self.feeding, errata)

    def test_las_celdas_con_dos_valores_no_estan_en_el_catalogo(self):
        for code in ("Disc SECT-I", "FS-1 VoltageD", "LoadB/NS FS-1", "LoadB/PP FS-1"):
            self.assertNotIn(code, self.feeding, code)

    def test_los_doce_codigos_de_la_leyenda_de_anclaje_estan_habilitados(self):
        # Los que dibuja el bloque ANCHORAGE de la leyenda de los workbooks.
        legend = {
            "CP+AnMC", "FP+AnMC", "CP/Tunnel", "FP/Tunnel",
            "CP/TX-P/1100", "CP/TX-T", "CP/TX-W/1100",
            "AnFW", "AnRW", "AnFW/Tunnel", "AnRW/Tunnel", "IO",
        }
        faltan = {c for c in legend if self.anchorage.get(c) != "SI"}
        self.assertEqual(faltan, set())

    def test_las_nueve_variantes_de_CP_TX_son_codigos_propios(self):
        """CP/TX-P, -T y -W, cada uno con o sin longitud: nueve anclajes distintos.

        La leyenda solo dibuja tres, pero los datos traen las nueve combinaciones y son
        anclajes diferentes, no un tipo con una medida anotada al lado.
        """
        for base in ("CP/TX-P", "CP/TX-T", "CP/TX-W"):
            for code in (base, base + "/1100", base + "/1350"):
                self.assertEqual(self.anchorage.get(code), "SI", code)

    def test_el_catalogo_de_anclaje_no_tiene_celdas_con_dos_codigos(self):
        """Ninguna fila del catalogo puede ser dos anclajes escritos seguidos.

        Entraron 28 asi —'CP+AnMC IO', 'AnRW AnRW', 'IO FP+AnMC'— porque el catalogo se
        cosecha leyendo cada celda de las hojas Track como si fuera un codigo. Una vez
        dentro, nada las distinguia de las de verdad y un perfil con dos anclajes acababa
        con uno inventado.
        """
        for code in ("CP+AnMC IO", "AnRW AnRW", "IO FP+AnMC", "CP+AnMC CP+AnMC",
                     "FP+AnMC AnRW", "IO AnRW", "CP+AnMC FP+AnMC", "AnFW AnFW",
                     "FP+AnMCAnRW", "AnRWAnRW"):
            self.assertNotIn(code, self.anchorage, code)

    def test_el_catalogo_de_anclaje_no_tiene_longitudes_ni_anotaciones(self):
        """La longitud de semitension y la via no son parte del anclaje.

        La leyenda las declara aparte ('SEMI TENSION LENGTH xxx m.'), y el dominio no
        tiene donde guardar la longitud: se queda el codigo y el numero se tira.
        """
        for code in ("CP+AnMC 265,00", "CP+AnMC CP+AnMC 527,00", "CP+AnMC 287",
                     "TRACK 5", "AnRW Portal", "AnRW2 Portal", "AnMP T1",
                     "AnRW2 track 1"):
            self.assertNotIn(code, self.anchorage, code)

    def test_las_erratas_de_tecleo_del_anclaje_no_son_codigos_nuevos(self):
        """Cambiar dos letras de sitio no crea un anclaje."""
        for errata in ("PF+AnMC", "CP+AnCM", "FP+AnCM", "CP+AMC", "CP+TX-P/1100"):
            self.assertNotIn(errata, self.anchorage, errata)

    def test_AnRW2_es_un_codigo_propio_y_no_dos_AnRW(self):
        """El numero forma parte del codigo. Son 108 celdas y un @ManyToMany no guarda
        cantidades: si fuera 'dos AnRW' el dos se perderia sin que se notara."""
        self.assertEqual(self.anchorage.get("AnRW2"), "SI")
        self.assertEqual(self.anchorage.get("AnRW"), "SI")
        self.assertEqual(self.anchorage.get("AnRW2/Tunnel"), "SI")

    def test_el_uno_de_AnRW1_es_la_cuenta_y_no_parte_del_codigo(self):
        """'AnRW1' es 'AnRW'. Que el 2 de 'AnRW2' SI cuente no lo contradice: uno de algo
        es ese algo, y el catalogo no puede llevar las dos grafias de lo mismo.

        Igual con el 2 delante ('2AnRW'), que ahi cuenta anclajes y no nombra otro: el
        modelo es un conjunto, asi que colapsa a uno.
        """
        self.assertNotIn("AnRW1", self.anchorage)
        self.assertNotIn("2AnRW", self.anchorage)
        self.assertNotIn("2AnRW Portal", self.anchorage)

    def test_una_celda_que_solo_es_una_anotacion_no_entra_al_catalogo(self):
        """'(Track 02)' y 'TRACK 5' no dejan codigo al quitarles lo que sobra.

        No es lo mismo que "no hay nada que partir", y tratarlo igual las dejaba dentro
        del catalogo con ENABLED=NO, como si fueran codigos pendientes de decidir.
        """
        for anotacion in ("(Track 02)", "TRACK 5"):
            self.assertNotIn(anotacion, self.anchorage, anotacion)

    def test_el_marcador_de_tipo_de_mensula_sin_declarar_existe_y_se_ve(self):
        """'UNKNOWN' no sale de ningun workbook: lo pone este proyecto.

        Existe porque el tipo de mensula es una relacion obligatoria y hay 217 huecos con
        medidas reales y sin tipo: sin un codigo al que colgarlos, esas mensulas se
        pierden enteras. Sale con ORIGEN=MTO justamente para que se distinga de un codigo
        del origen y se pueda auditar.
        """
        enabled, origen = self.cantilever_types.get("UNKNOWN", (None, None))
        self.assertEqual(enabled, "SI")
        self.assertEqual(origen, "MTO")

    def test_la_mensula_de_catenaria_rigida_ya_estaba_en_el_catalogo(self):
        """'OCR' (Overhead Conductor Rail) lo escribe EP9B en la columna del tipo.

        Por eso los huecos de EP6 con 'OCR SUPPORT' no necesitan un codigo nuevo: es el
        mismo concepto, y darle otro nombre partiria en dos una sola cosa.
        """
        enabled, origen = self.cantilever_types.get("OCR", (None, None))
        self.assertEqual(enabled, "SI")
        self.assertIn("TRACK", str(origen))

    def test_el_catalogo_de_brazos_son_solo_los_tipos(self):
        """La longitud del brazo no es parte de su tipo.

        El origen escribe 'PHQ-1150' —el tipo y su longitud juntos— y el catalogo se
        cosecha leyendo la celda entera, asi que llego a tener 71 codigos de los cuales
        62 eran un tipo con su medida dentro del nombre. Ninguno se usaba, porque el
        maestro de perfiles ya los separa, pero estaban marcados "pendiente de decidir":
        habilitar uno habria guardado la misma medida dos veces, en steady_arm.length y
        dentro del codigo.
        """
        base = {"BC", "BCE", "BS", "BTC", "PH", "PH-C", "PH-Q", "PHC", "PHQ"}
        habilitados = {c for c, enabled in self.steady_arm_types.items() if enabled == "SI"}
        self.assertEqual(habilitados, base)
        for compuesto in ("PHQ-1150", "PH-1150", "PHQ-950", "PHC-1150", "BTC-1651",
                          "PH950", "PHQ- 1150", "PHC-1500E"):
            self.assertNotIn(compuesto, self.steady_arm_types, compuesto)

    def test_no_queda_ningun_tipo_de_brazo_por_decidir(self):
        """Los dos que salieron nombrados ya estan resueltos, y de dos formas distintas.

        'BS' era un tipo de verdad —una sola hoja lo usa, EP9B / HR Track 2 LOD_S— y se
        da de alta. 'BHC' era 'PHC' mal tecleado y se corrige con un alias, no con un
        codigo nuevo: dar de alta una errata la convierte en un tipo para siempre.

        Que ninguno de los dos siga como fila del catalogo es lo que prueba que se han
        resuelto de verdad y no aparcado.
        """
        self.assertEqual({c for c, e in self.steady_arm_types.items() if e != "SI"}, set())
        for errata in ("BS-1400", "BHC-1150", "BHC"):
            self.assertNotIn(errata, self.steady_arm_types, errata)

    def test_el_soporte_de_catenaria_rigida_esta_habilitado(self):
        """'OCR SUPPORT' ya se usa para deducir el tipo de mensula de EP6 TSA-THA.

        Si el dato decide lo que se carga, guardarlo es lo minimo coherente: dejarlo sin
        habilitar significaria que influye en la importacion y no llega a la base de datos.
        """
        self.assertEqual(self.support_types.get("OCR SUPPORT"), "SI")

    def test_las_grafias_de_la_suspension_aislada_son_una_sola(self):
        """Repartidas, cada variante quedaba por debajo del umbral de atencion."""
        for grafia in ("MP_Isusp", "MW/Isusp", "MWISusp", "MW/CW_ISusp", "MW/MC-ISusp",
                       "S1B7", "S/B7"):
            self.assertNotIn(grafia, self.support_types, grafia)

    def test_los_aparatos_de_alimentacion_no_son_soportes(self):
        """'SECT-I', 'FS1' y 'FS/PP2' son de la columna 'Sectioning Feeding'.

        Fuera del catalogo de soportes: asi la celda sale nombrada con su hoja y su fila,
        que es lo que hay que corregir, en vez de inventar tres tipos de soporte.
        """
        for fuga in ("SECT-I", "FS1", "FS/PP2"):
            self.assertNotIn(fuga, self.support_types, fuga)

    def test_unique_solution_es_un_valor_real_y_no_un_marcador(self):
        """No es un hueco: es la cimentacion que necesita solucion a medida, aparte, porque
        ninguna convencional sirve. Estuvo en code_rejections por una suposicion mia y eso
        dejaba 468 perfiles sin cimentacion. Su tipo es USX justamente porque no encaja en
        ninguna de las seis familias estructurales."""
        self.assertEqual(self.foundation.get("UNIQUE SOLUTION"), "SI")
        self.assertEqual(self.foundation_types.get("USX"), "SI")
        for grafia in ("Unique Solution", "U.S.", "U.SOLUTION", "U. SOLUTION"):
            self.assertNotIn(grafia, self.foundation, grafia)

    def test_las_erratas_de_t_sign_found_se_funden_con_el_codigo_bueno(self):
        for errata in ("T-SING FOUND.", "T-SIGN FOUND,"):
            self.assertNotIn(errata, self.foundation, errata)
        self.assertEqual(self.foundation.get("T-SIGN FOUND."), "SI")

    def test_no_queda_nada_sin_reconocer(self):
        # Si esta hoja trae filas, el catalogo sale incompleto y el generador
        # termina con codigo != 0.
        self.assertEqual(self.unknown_rows, 0)


if __name__ == "__main__":
    unittest.main()
