"""Reglas de lectura del plano de seccionamiento que no necesitan el DXF.

Lo que se prueba aqui es lo que decide un dato del libro de revision a partir de un texto o de
unas coordenadas: como se lee un KP, cuantos aparatos nombra un rotulo, que es un rotulo de
subestacion y no un seccionador, cuando un simbolo esta en la columna de su rotulo y como se
elige la via del maestro. La lectura del DXF y la asignacion (ezdxf, scipy) quedan fuera: CI
instala solo openpyxl y pyyaml, y el modulo las importa dentro de las funciones que las usan.

    python3 -m unittest discover -s data/tools/tests
"""

import importlib.util
import os
import unittest

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def _load(name):
    spec = importlib.util.spec_from_file_location(name, os.path.join(BASE, name + ".py"))
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


bsr = _load("build_sectioning_review")


class KpTest(unittest.TestCase):

    def test_the_plan_writes_km_plus_metres_and_the_master_wants_metres(self):
        self.assertEqual(bsr.parse_kp("KP93+451"), ("KP93+451", 93451))
        self.assertEqual(bsr.parse_kp("KP 13+523"), ("KP13+523", 13523))
        self.assertEqual(bsr.parse_kp("K.P.\t1+915"), ("KP1+915", 1915))
        self.assertEqual(bsr.parse_kp("(KP101+200)"), ("KP101+200", 101200))

    def test_a_bare_kp_and_a_decimal_kp_are_read_too(self):
        self.assertEqual(bsr.parse_kp("131+200"), ("KP131+200", 131200))
        self.assertEqual(bsr.parse_kp("KP9+320.5"), ("KP9+320.5", 9320.5))

    def test_an_x_in_the_kp_is_an_unknown_kp_not_a_zero(self):
        self.assertEqual(bsr.parse_kp("KPXX+XXX"), ("KPXX+XXX", None))
        self.assertEqual(bsr.parse_kp("K.P 2+XXX"), ("KP2+XXX", None))

    def test_a_name_is_not_a_kp(self):
        self.assertEqual(bsr.parse_kp("TSA-BF06"), (None, None))


class LabelTest(unittest.TestCase):

    def test_a_disconnector_label_is_a_name_and_a_kp(self):
        label = bsr.parse_label("TSA-BF06\nKP93+451", "0_Disconnectors")
        self.assertEqual(label["kind"], "device")
        self.assertEqual(label["names"], ["TSA-BF06"])
        self.assertEqual(label["kp_m"], 93451)

    def test_a_slash_names_two_devices_of_the_same_installation(self):
        self.assertEqual(
            bsr.parse_label("UN2-NS3/NS4\nKP86+554", "0_Disconnectors")["names"],
            ["UN2-NS3", "UN2-NS4"],
        )
        self.assertEqual(bsr.expand_names("BSD-09/BSD-TE09"), ["BSD-09", "BSD-TE09"])

    def test_a_comma_list_replaces_the_tail_of_the_first_name(self):
        self.assertEqual(
            bsr.expand_names("DP2-TE105B, 104B, 103B"), ["DP2-TE105B", "DP2-TE104B", "DP2-TE103B"]
        )

    def test_a_space_after_the_dash_is_a_typo(self):
        self.assertEqual(
            bsr.parse_label("NRY- 03\nKP29+XXX", "0_Disconnectors")["names"], ["NRY-03"]
        )

    def test_name_and_kp_on_the_same_line(self):
        label = bsr.parse_label("HSA-FP1.1 KP9+316", "0_Disconnectors")
        self.assertEqual((label["names"], label["kp_m"]), (["HSA-FP1.1"], 9316))

    def test_neutral_sections_and_substations_are_context_not_devices(self):
        self.assertEqual(
            bsr.parse_label("NS ANAVA\nKP25+131", "0_Disconnectors")["kind"], "context"
        )
        self.assertEqual(bsr.parse_label("TS DOR\nKP29+700", "0_Disconnectors")["kind"], "context")
        self.assertEqual(bsr.parse_label("BY-PASS FEEDER FOR EMF", "0_Disconnectors")["names"], [])

    def test_feeders_inside_a_substation_are_not_matched_to_disconnectors(self):
        # 'F1.1' apunta a geometria gris dentro del recuadro de la subestacion: si se tratase como
        # un seccionador, la asignacion le daria el simbolo de catenaria de al lado.
        self.assertEqual(bsr.parse_label("F1.1", "0_Disconnectors")["kind"], "feeder")

    def test_depot_names_only_count_on_the_disconnector_layer(self):
        # Los del deposito de Haifa sin guion ('LB121'); con guion ('LE2211-9') se leen como
        # cualquier seccionador, y su posicion en la zona rayada ya los deja fuera.
        self.assertEqual(bsr.parse_label("LB121", "0_Disconnectors")["kind"], "depot")
        # 'L4' en la capa del deposito de Lod es una via, no un seccionador.
        self.assertEqual(bsr.parse_label("L4", "Lod_Depot")["names"], [])

    def test_station_code_from_the_name(self):
        self.assertEqual(bsr.station_code("TSA-BF06"), "TSA")
        self.assertEqual(bsr.station_code("ASKD-02BIS"), "ASK-D")
        self.assertEqual(bsr.station_code("DLOD-31"), "LOD_D")
        self.assertEqual(bsr.station_code("*TN1-TE1"), "TN1")
        self.assertEqual(bsr.station_code("F1.1"), "")


class FunctionTest(unittest.TestCase):

    def test_what_the_name_says(self):
        self.assertEqual(bsr.function_from_name("KAF-NS2"), "NS")
        self.assertEqual(bsr.function_from_name("TSA-BF06"), "BF")
        self.assertEqual(bsr.function_from_name("NIZ-B01"), "B")
        self.assertEqual(bsr.function_from_name("HFT-FP1.2"), "FP")
        self.assertEqual(bsr.function_from_name("TN3-TE1"), "TE")
        self.assertEqual(bsr.function_from_name("NIZ-04"), "")

    def test_proposal_when_there_is_no_pole_to_copy_from(self):
        self.assertEqual(bsr.function_proposal("SECCIONADOR", True, "", "IO"), "LoadB/IO")
        self.assertEqual(bsr.function_proposal("SECCIONADOR", False, "", "IO"), "Disc/IO")
        self.assertEqual(bsr.function_proposal("SECCIONADOR", True, "NS", "IO"), "LoadB/NS")
        self.assertEqual(bsr.function_proposal("SECCIONADOR", False, "", "SI"), "Disc/SI")
        # Los de alimentacion, del portico de la subestacion: el catalogo no tiene un FP.
        self.assertEqual(bsr.function_proposal("SECCIONADOR", True, "FP", ""), "LoadB")
        self.assertEqual(bsr.function_proposal("SECCIONADOR", False, "FP", ""), "Disc")
        self.assertEqual(
            bsr.function_proposal("SECCIONADOR DE PUESTA A TIERRA", False, "", ""), "ED"
        )
        # Sin pista no se inventa nada.
        self.assertEqual(bsr.function_proposal("SECCIONADOR", True, "", ""), "")


class GeometryTest(unittest.TestCase):

    # Un rotulo vertical (girado 90 grados) de 2 x 30 unidades, insertado en (100, 0).
    VERTICAL = (90.0, 100.0, 0.0, 0.0, 30.0, -2.0, 0.0)

    def test_a_symbol_below_a_vertical_label_is_in_its_column(self):
        du, dv = bsr.outside(self.VERTICAL, 101.0, -12.0)
        self.assertAlmostEqual(dv, 0.0)
        self.assertAlmostEqual(du, 12.0)

    def test_a_symbol_beside_the_label_is_off_the_column(self):
        du, dv = bsr.outside(self.VERTICAL, 106.0, 10.0)
        self.assertAlmostEqual(du, 0.0)
        self.assertAlmostEqual(dv, 4.0)

    def test_directions_not_senses(self):
        self.assertAlmostEqual(bsr.angle_diff(0, 180), 0.0)
        self.assertAlmostEqual(bsr.angle_diff(10, 170), 20.0)

    def test_kp_between_two_neighbours(self):
        self.assertAlmostEqual(bsr.interpolate_kp(5.0, [(0.0, 1000.0), (10.0, 1100.0)]), 1050.0)
        self.assertEqual(bsr.interpolate_kp(5.0, [(0.0, 1000.0)]), 1000.0)
        self.assertIsNone(bsr.interpolate_kp(5.0, []))


class MasterTrackTest(unittest.TestCase):

    TRACKS = [
        {"name": "TRACK 1", "num": "1", "stations": ["ZIC", "BIN", "HAD", "KFA"]},
        {"name": "TRACK 3 BIN", "num": "3", "stations": ["BIN"]},
        {"name": "TRACK 3 HAD", "num": "3", "stations": ["HAD"]},
        {"name": "TRACK 3 KFA", "num": "3", "stations": ["KFA"]},
    ]

    def test_the_number_of_the_master_track_name(self):
        self.assertEqual(bsr.track_number_of("TRACK 3 BIN"), "3")
        self.assertEqual(bsr.track_number_of("TRACK 04 KFA"), "4")
        self.assertIsNone(bsr.track_number_of("TRACK INT NORTH"))
        self.assertIsNone(bsr.track_number_of("TRACK 1.2 Ashdod - Nizanim - Ashkelon"))

    def test_a_unique_number_needs_no_station(self):
        self.assertEqual(bsr.master_track(self.TRACKS, "", "1"), ("TRACK 1", ""))

    def test_the_station_breaks_the_tie(self):
        self.assertEqual(bsr.master_track(self.TRACKS, "HAD", "3"), ("TRACK 3 HAD", ""))

    def test_no_station_no_choice(self):
        name, why = bsr.master_track(self.TRACKS, "", "3")
        self.assertEqual(name, "")
        self.assertIn("ambigua", why)

    def test_a_number_the_master_does_not_have(self):
        self.assertEqual(
            bsr.master_track(self.TRACKS, "BIN", "45"), ("", "la via 45 no existe en el maestro")
        )


class ConfidenceTest(unittest.TestCase):

    def test_a_leader_line_is_certain(self):
        self.assertEqual(bsr.confidence({"how": "leader"}), "A")

    def test_aligned_close_and_without_rival(self):
        self.assertEqual(
            bsr.confidence({"how": "aligned", "cost": 10, "du": 10, "dv": 0, "alt": 40}), "A"
        )

    def test_a_rival_almost_as_close_lowers_it(self):
        self.assertEqual(
            bsr.confidence({"how": "aligned", "cost": 10, "du": 10, "dv": 0, "alt": 12}), "B"
        )

    def test_the_relaxed_pass_is_always_doubtful(self):
        self.assertEqual(bsr.confidence({"how": "relaxed"}), "C")


class DrawingFieldsTest(unittest.TestCase):

    def test_the_drawn_blade_is_the_normal_state(self):
        self.assertEqual(bsr.normally_open("ABIERTO"), "SI")
        self.assertEqual(bsr.normally_open("CERRADO"), "NO")
        self.assertEqual(bsr.normally_open(""), "")

    def test_the_drive_circle_is_the_motor(self):
        self.assertEqual(bsr.drive_type(True), "MOTOR")
        self.assertEqual(bsr.drive_type(False), "")

    def test_only_the_haifa_depot_is_out_of_every_master_ep(self):
        self.assertTrue(bsr.future_ep(bsr.HAIFA_STATUS))
        for status in [bsr.EN_SERVICE] + list(bsr.STATUS_BY_LAYER.values()):
            self.assertFalse(bsr.future_ep(status), status)


class PoleCostTest(unittest.TestCase):
    """Que poste puede ser el de un seccionador, y cuanto cuesta casarlos."""

    def device(self, **overrides):
        r = {
            "kp_m": 1000,
            "tracks": ["1"],
            "type": "SECCIONADOR",
            "name_function": "",
            "on_load": True,
        }
        r.update(overrides)
        return r

    def pole(self, *codes, kp=1002.0, track="1"):
        return {"kp": kp, "track": track, "dcodes": list(codes)}

    def test_the_kp_difference_with_penalties_for_track_and_load(self):
        self.assertEqual(bsr.pole_cost(self.device(), self.pole("LoadB/IO")), 2.0)
        self.assertEqual(bsr.pole_cost(self.device(), self.pole("LoadB/IO", track="2")), 62.0)
        self.assertEqual(bsr.pole_cost(self.device(), self.pole("Disc/IO")), 27.0)
        self.assertIsNone(bsr.pole_cost(self.device(), self.pole("LoadB/IO", kp=1081.0)))

    def test_an_earthing_disconnector_only_goes_on_a_pole_marked_ed(self):
        earthing = self.device(type="SECCIONADOR DE PUESTA A TIERRA", on_load=False)
        self.assertEqual(bsr.pole_cost(earthing, self.pole("ED/T")), 2.0)
        self.assertEqual(bsr.pole_cost(earthing, self.pole("LoadB/ED")), 2.0)
        self.assertIsNone(bsr.pole_cost(earthing, self.pole("Disc/IO")))
        # TE en el nombre cuenta igual que el bloque de puesta a tierra.
        self.assertIsNone(bsr.pole_cost(self.device(name_function="TE"), self.pole("LoadB/NS")))

    def test_a_line_disconnector_does_not_take_an_earthing_pole(self):
        self.assertIsNone(bsr.pole_cost(self.device(), self.pole("ED")))
        self.assertEqual(bsr.pole_cost(self.device(), self.pole("LoadB/ED")), 2.0)

    def test_the_two_kinds_that_are_not_on_a_line_pole(self):
        self.assertTrue(bsr.on_portal(self.device(name_function="FP")))
        self.assertFalse(bsr.on_portal(self.device(name_function="NS")))
        self.assertTrue(bsr.earthing(self.device(name_function="TE")))
        self.assertTrue(bsr.earthing(self.device(type="SECCIONADOR DE PUESTA A TIERRA")))


class DisconnectorRowsTest(unittest.TestCase):
    """Lo que decide si una fila de DISCONNECTORS entra y si sale lista para cargar."""

    POLE = {
        "ep": "EP6",
        "via": "TRACK 1 TLV SAVIDOR",
        "profile": "93-1.05",
        "kp": 93453.0,
        "track": "1",
        "dcodes": ["Disc/IO"],
    }
    MASTER = {
        "tracks": {"EP6": [{"name": "TRACK 1 TLV SAVIDOR", "num": "1", "stations": ["TSA"]}]},
        "profiles": [POLE],
    }

    def record(self, **overrides):
        r = {
            "handle": "H1",
            "layer": "0_Disconnectors",
            "status": bsr.EN_SERVICE,
            "x": 0.0,
            "y": 0.0,
            "name": "TSA-06",
            "label_text": "TSA-06 | KP93+451",
            "kp_txt": "KP93+451",
            "kp_m": 93451,
            "names_in_label": 1,
            "how": "leader",
            "confidence": "A",
            "type": "SECCIONADOR",
            "on_load": False,
            "state": "CERRADO",
            "has_drive": True,
            "vd": False,
            "prefix": "TSA",
            "plan_station": "",
            "tracks": ["1"],
            "sections": ["0_Has1"],
            "bridged": "",
            "name_function": "",
            "ep": "EP6",
            "ep_from": "prefijo",
            "station": "TSA",
            "station_from": "prefijo",
            "pole": self.POLE,
            "pole_dkp": 2.0,
        }
        r.update(overrides)
        return r

    def row(self, master=None, **overrides):
        rows, outside = bsr.disconnector_rows([self.record(**overrides)], master or self.MASTER)
        self.assertEqual(len(rows), 1, outside)
        return rows[0]

    def test_a_disconnector_on_its_pole_is_ready(self):
        row = self.row()
        self.assertEqual(
            (row["PROFILE_ID"], row["DISCONNECTOR_FUNCTION"], row["ENABLED"], row["REVISAR"]),
            ("93-1.05", "Disc/IO", "SI", "NO"),
        )
        self.assertEqual((row["NORMALLY_OPEN"], row["DRIVE_TYPE"]), ("NO", "MOTOR"))

    def test_what_is_drawn_as_future_or_disabled_enters_if_it_is_in_a_master_ep(self):
        # Holtz: la zona neutra esta dibujada como deshabilitada, pero esta en EP9A.
        for status in ("FUTURO", "DESHABILITADO (NS)", "VIA NO ELECTRIFICADA"):
            row = self.row(status=status)
            self.assertEqual(row["ESTADO_DIBUJO"], status)
            self.assertIn("dibujado como " + status, row["MOTIVO_REVISAR"])
            self.assertEqual(row["ENABLED"], "SI", "el estado de dibujo es una nota")

    def test_what_is_not_in_a_master_ep_goes_outside_and_haifa_nowhere(self):
        rows, outside = bsr.disconnector_rows([self.record(ep=None, pole=None)], self.MASTER)
        self.assertEqual((len(rows), len(outside)), (0, 1))
        rows, outside = bsr.disconnector_rows(
            [self.record(ep=None, pole=None, status=bsr.HAIFA_STATUS)], self.MASTER
        )
        self.assertEqual((rows, outside), ([], []))

    def test_without_a_pole_the_row_waits_for_a_person(self):
        # El poste es opcional, pero el generador no distingue un seccionador que no esta en un
        # poste de un poste que no ha encontrado.
        row = self.row(pole=None, bridged="IO")
        self.assertEqual((row["PROFILE_ID"], row["ENABLED"]), ("", "NO"))
        self.assertIn("deja PROFILE_ID vacio y pon ENABLED = SI", row["MOTIVO_REVISAR"])

    def test_feeder_and_earthing_disconnectors_without_a_pole_are_a_note(self):
        # Los de los porticos de subestacion y los de puesta a tierra no van en un poste: que no
        # lo tengan no es lo que deja la fila esperando.
        def blocking(row):
            return [x for x in row["MOTIVO_REVISAR"].split("; ") if x.startswith(bsr.BLOCKING)]

        feeder = self.row(pole=None, name="HSA-FP1.1", name_function="FP")
        self.assertEqual((feeder["PROFILE_ID"], feeder["DISCONNECTOR_FUNCTION"]), ("", "Disc"))
        self.assertIn("sin poste por ser de alimentacion", feeder["MOTIVO_REVISAR"])
        self.assertEqual(blocking(feeder), ["funcion propuesta desde el plano (Disc)"])
        earthing = self.row(
            pole=None, name="TN3-TE3", type="SECCIONADOR DE PUESTA A TIERRA", name_function="TE"
        )
        self.assertIn("sin poste por ser de puesta a tierra", earthing["MOTIVO_REVISAR"])
        self.assertEqual(blocking(earthing), ["funcion propuesta desde el plano (ED)"])

    def test_the_own_kp_only_without_a_pole(self):
        # El KP propio es solo de un seccionador sin poste: con poste, es el del perfil (V26).
        self.assertEqual(self.row()["KP"], "")
        # KP_POSTE es el del poste, que distingue los dos de una via que repite su PROFILE_ID.
        self.assertEqual(self.row()["KP_POSTE"], self.POLE["kp"])
        self.assertEqual(self.row(pole=None, name_function="FP")["KP_POSTE"], "")
        self.assertEqual(self.row(pole=None, name_function="FP")["KP"], 93451)
        self.assertEqual(self.row(pole=None, name_function="FP", kp_m=None, kp_txt=None)["KP"], "")

    def test_an_unknown_normal_state_is_a_note(self):
        row = self.row(state="")
        self.assertEqual((row["NORMALLY_OPEN"], row["ENABLED"]), ("", "SI"))
        self.assertIn("estado normal sin determinar", row["MOTIVO_REVISAR"])

    def test_another_track_is_a_note_only_while_the_kp_holds(self):
        # El plano numera las vias de la estacion, no las del maestro: con el KP casado, una via
        # distinta no impide cargar. Sin otro poste al mismo KP, que seria otra cosa (abajo).
        other = dict(self.POLE, track="3", via="TRACK 3 TLV SAVIDOR")
        master = dict(self.MASTER, profiles=[other])
        close = self.row(master, pole=other, pole_dkp=2.0)
        self.assertEqual(close["ENABLED"], "SI")
        self.assertIn("via distinta", close["MOTIVO_REVISAR"])
        far = self.row(master, pole=other, pole_dkp=20.0)
        self.assertEqual(far["ENABLED"], "NO")
        self.assertIn("el DXF dice via 1", far["MOTIVO_REVISAR"])

    def test_the_poles_near_the_kp_are_the_options_of_profile_id_and_kp_poste(self):
        # Los del EP a menos de 80 m del KP del rotulo, del mas cercano al mas lejano: el de 149 m
        # y el de otro EP no se ofrecen.
        near = dict(self.POLE, profile="93-1.04", kp=93430.0, dcodes=[])
        far = dict(self.POLE, profile="93-1.08", kp=93600.0)
        other_ep = dict(self.POLE, ep="EP5", profile="93-1.05", kp=93451.0)
        master = dict(self.MASTER, profiles=[far, near, self.POLE, other_ep])

        row = self.row(master=master)
        self.assertEqual(row["_options"]["PROFILE_ID"], ["93-1.05", "93-1.04"])
        self.assertEqual(row["_options"]["KP_POSTE"], [93453.0, 93430.0])
        self.assertEqual(
            row["POSTES_CERCANOS"],
            "93-1.05 (TRACK 1 TLV SAVIDOR, kp 93453, a 2 m, Disc/IO) | "
            "93-1.04 (TRACK 1 TLV SAVIDOR, kp 93430, a 21 m)",
        )
        # El poste de la fila va primero aunque no sea el mas cercano.
        own = dict(near, dcodes=["Disc/IO"])
        self.assertEqual(
            self.row(master=master, pole=own)["_options"]["PROFILE_ID"], ["93-1.04", "93-1.05"]
        )
        # Sin KP en el rotulo no hay a que distancia medir.
        self.assertEqual(self.row(master=master, kp_m=None)["_options"], {})

    def test_the_pole_of_the_row_heads_both_lists_even_beyond_the_nearby_ones(self):
        # El poste casado puede estar a mas de 80 m del KP del rotulo: su KP tambien va primero, o
        # el KP_POSTE de la fila no estaria entre las opciones de su propia celda.
        distant = dict(self.POLE, profile="93-1.20", kp=93600.0)
        master = dict(self.MASTER, profiles=[self.POLE, distant])
        row = self.row(master=master, pole=distant, pole_dkp=149.0)
        self.assertEqual(row["_options"]["PROFILE_ID"], ["93-1.20", "93-1.05"])
        self.assertEqual(row["_options"]["KP_POSTE"], [93600.0, 93453.0])

    def test_a_proposed_pole_brings_its_kp_and_is_not_a_twin(self):
        # Aceptar la propuesta no deja KP_POSTE a medias: en una via que repite el identificador
        # es lo que dice cual de los dos es.
        bare = dict(self.POLE, profile="93-1.04", kp=93440.0, dcodes=[], codes=[])
        master = dict(self.MASTER, profiles=[self.POLE, bare])
        proposed = self.row(master=master, pole=None, nearest_pole=bare, nearest_pole_d=11.0)
        self.assertEqual((proposed["PROFILE_ID"], proposed["KP_POSTE"]), ("93-1.04", 93440.0))
        self.assertEqual(proposed["_proposals"], {"VIA", "PROFILE_ID", "KP_POSTE"})
        self.assertEqual(proposed["_options"]["KP_POSTE"][0], 93440.0)
        # Un poste propuesto que coincide con uno casado no marca al casado: ya espera a alguien.
        records = [
            self.record(),
            self.record(
                handle="H2",
                name="TSA-07",
                pole=None,
                nearest_pole=dict(self.POLE, codes=[]),
                nearest_pole_d=2.0,
            ),
        ]
        rows, _ = bsr.disconnector_rows(records, master)
        ready = next(d for d in rows if d["HANDLE"] == "H1")
        self.assertEqual(ready["ENABLED"], "SI")
        self.assertNotIn("repite el poste", ready["MOTIVO_REVISAR"])

    def test_twin_poles_the_plan_track_does_not_settle_are_doubtful(self):
        int_pole = dict(self.POLE, via="TRACK INT SOUTH", profile="2-INT.13", track=None)
        ext_pole = dict(self.POLE, via="TRACK EXT SOUTH", profile="2-EXT.13", track=None)
        master = dict(self.MASTER, profiles=[int_pole, ext_pole])
        rows, _ = bsr.disconnector_rows([self.record(pole=int_pole)], master)
        self.assertEqual(rows[0]["ENABLED"], "NO")
        self.assertIn("via sin comprobar", rows[0]["MOTIVO_REVISAR"])
        self.assertIn("2-EXT.13", rows[0]["MOTIVO_REVISAR"])


class DropDownTest(unittest.TestCase):
    """Los desplegables del libro: las opciones de las celdas amarillas y naranjas."""

    def test_the_own_options_of_a_row_skip_blanks_and_repeats(self):
        row = {"_options": {"PROFILE_ID": ["93-1.05", "", "93-1.05", None], "NOMBRE": [""]}}
        self.assertEqual(bsr.own_options(row), {"PROFILE_ID": ["93-1.05"]})
        self.assertEqual(bsr.own_options({}), {})

    def test_the_loose_labels_beside_an_unnamed_symbol_are_its_name_options(self):
        rows = [
            {"NOMBRE": "", "X": 100.0, "Y": 0.0},
            {"NOMBRE": "LOD-03", "X": 100.0, "Y": 0.0},
            {"NOMBRE": "", "X": 5000.0, "Y": 0.0},
        ]
        orphans = [
            {"NOMBRE": "LOD-15", "X": 220.0, "Y": 0.0},
            {"NOMBRE": "LOD-14", "X": 150.0, "Y": 0.0},
            {"NOMBRE": "HFT-FP1.2", "X": 1600.0, "Y": 0.0},
        ]
        bsr.orphan_name_options(rows, orphans)
        self.assertEqual(rows[0]["_options"]["NOMBRE"], ["LOD-14", "LOD-15"])
        self.assertNotIn("_options", rows[1], "un simbolo con nombre no lo necesita")
        self.assertNotIn("_options", rows[2], "ningun rotulo suelto cerca")

    def test_the_function_catalogue_offers_the_disconnector_codes_first(self):
        try:
            import openpyxl
        except ImportError:
            raise unittest.SkipTest("openpyxl no instalado")
        import tempfile

        path = os.path.join(tempfile.mkdtemp(), "lov-master.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "LOVS"
        ws.append(["ENTIDAD", "CODIGO", "DESCRIPCION_ES", "ENABLED"])
        for entity, code, enabled in (
            ("DisconnectorFunction", "SurgeA", "SI"),
            ("DisconnectorFunction", "LoadB/NS", "SI"),
            ("DisconnectorFunction", "Disc/X", "NO"),
            ("DisconnectorFunction", "ED", "SI"),
            ("DisconnectorFunction", "Disc/IO", "SI"),
            ("Sectioning", "A/S", "SI"),
        ):
            ws.append([entity, code, None, enabled])
        wb.save(path)

        self.assertEqual(
            bsr.read_disconnector_functions(path), ["Disc/IO", "ED", "LoadB/NS", "SurgeA"]
        )
        # Sin el maestro de catalogos, los codigos que ya usan los postes.
        profiles = [{"codes": ["FS-1", "Disc/IO"]}, {"codes": ["LoadB"]}]
        self.assertEqual(
            bsr.read_disconnector_functions(
                os.path.join(os.path.dirname(path), "no.xlsx"), profiles
            ),
            ["Disc/IO", "LoadB", "FS-1"],
        )

    def test_the_workbook_carries_the_dropdowns(self):
        try:
            import openpyxl
        except ImportError:
            raise unittest.SkipTest("openpyxl no instalado")
        import tempfile

        master = {
            "stations": {"EP6": {"TSA", "HER"}},
            "tracks": {"EP6": [{"name": "TRACK 1 TLV SAVIDOR"}, {"name": "TRACK 2 TLV SAVIDOR"}]},
        }
        choices = {
            "lists": bsr.review_lists(master, ["Disc/IO", "LoadB"]),
            "columns": bsr.review_choices(),
        }
        disconnectors = [
            {
                "EP": "EP6",
                "ESTACION": "TSA",
                "PROFILE_ID": "93-1.05",
                "KP_POSTE": 93453.41,
                "ENABLED": "SI",
                "HANDLE": "H1",
                "_options": {"PROFILE_ID": ["93-1.05", "93-1.04"], "KP_POSTE": [93453.41, 93430]},
            }
        ]
        path = os.path.join(tempfile.mkdtemp(), "review.xlsx")
        bsr.write_review(path, self.sheets(disconnectors), [], choices)

        wb = openpyxl.load_workbook(path)
        options = wb["OPCIONES"]
        header = [c.value for c in options[1]]
        stations = header.index("ESTACIONES EP6") + 1
        self.assertEqual(
            [options.cell(row=i, column=stations).value for i in (2, 3)], ["HER", "TSA"]
        )
        validations = wb["DISCONNECTORS"].data_validations.dataValidation

        def of(cell):
            found = [dv for dv in validations if cell in dv.sqref]
            self.assertEqual(len(found), 1, "una sola validacion por celda: %s" % cell)
            return found[0]

        letter = openpyxl.utils.get_column_letter(stations)
        self.assertEqual(of("B2").formula1, "OPCIONES!$%s$2:$%s$3" % (letter, letter))
        self.assertEqual(of("B2").errorStyle, "warning", "una estacion de otro EP, con aviso")
        # Las opciones propias de la fila, en su fila de la hoja oculta y con su tipo: un KP
        # escrito como texto no seria un numero en un Excel con coma decimal.
        self.assertEqual(of("D2").formula1, "OPCIONES_FILA!$D$2:$E$2")
        self.assertEqual(of("D2").errorStyle, "warning")
        self.assertEqual(of("E2").formula1, "OPCIONES_FILA!$D$3:$E$3")
        own = wb["OPCIONES_FILA"]
        self.assertEqual(own.sheet_state, "hidden")
        self.assertEqual(
            [[c.value for c in row] for row in own.iter_rows(min_row=2)],
            [
                ["DISCONNECTORS", 2, "PROFILE_ID", "93-1.05", "93-1.04"],
                ["DISCONNECTORS", 2, "KP_POSTE", 93453.41, 93430],
            ],
        )
        self.assertIn("OPCIONES!$A$", of("F2").formula1, "ENABLED: SI_NO, la primera lista")
        self.assertIn(of("F2").errorStyle, (None, "stop"))

    @staticmethod
    def sheets(disconnectors):
        """Las hojas que RESUMEN cuenta, vacias salvo DISCONNECTORS."""

        def sheet(title, columns, data=()):
            return title, columns, 0, set(), list(data), None

        return [
            sheet(
                "DISCONNECTORS",
                [
                    "EP",
                    "ESTACION",
                    "VIA",
                    "PROFILE_ID",
                    "KP_POSTE",
                    "ENABLED",
                    "REVISAR",
                    "ESTADO_DIBUJO",
                    "HANDLE",
                ],
                disconnectors,
            ),
            sheet("SECCIONADORES_DXF", ["HANDLE", "ESTADO_DIBUJO"]),
            sheet("ESTACION_POR_PREFIJO", ["PREFIJO", "EP", "ESTACION_CORRECTA"]),
            sheet("SECCIONADORES_FUERA_MAESTRO", ["HANDLE"]),
            sheet("AISLADORES_DXF", ["HANDLE"]),
            sheet("SECTION_INSULATORS", ["HANDLE", "KP", "ENABLED"]),
            sheet("AISLADORES_FUERA_MAESTRO", ["HANDLE"]),
            sheet("SECTION_INSULATOR_SWITCHES", ["HANDLE_AISLADOR", "TANGENTE"]),
            sheet("ROTULOS_SIN_SIMBOLO", ["ROTULO"]),
            sheet("POSTES_SIN_CASAR", ["PROFILE_ID"]),
        ]


if __name__ == "__main__":
    unittest.main()
