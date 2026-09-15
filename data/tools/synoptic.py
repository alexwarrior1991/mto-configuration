#!/usr/bin/env python3
"""Lector del formato SINOPTICO: una hoja con las dos vias en espejo.

Los once workbooks de los EPs ferroviarios traen una hoja 'HR Track' por via, con una
cabecera y las mismas ~36 columnas logicas. El sinoptico de la Linha Rubi (Metro do
Porto, RUBI.xlsx) no se parece en nada: es UNA hoja con las dos vias en ESPEJO, la via 2
a la izquierda (columnas B..AA, leyendo hacia fuera) y la via 1 a la derecha (AP..BM),
separadas por una banda central de columnas sin cabecera. Las cabeceras estan en
portugues y en espanol, hay dos columnas que solo se distinguen por una tilde
('Tipologia' es el sistema de catenaria; 'Tipología' es la funcion del apoyo), y las
medidas de la segunda mensula van en columnas SIN cabecera pegadas a las de la primera.

Este modulo es lo que hace falta para leer esa hoja y no tiene ninguna dependencia de
los generadores: lo usan LOS DOS —build_lov_master.py cosecha con el los codigos y
build_profile_master.py los perfiles—, igual que workbook_common.py, para que el
catalogo y las referencias de los perfiles no puedan discrepar sobre lo que dice una
celda. Todo lo que este modulo decide viene de la seccion 'synoptic' de aliases.yml
y de 'format: synoptic' en topology.yml.

Reglas de la casa que se conservan:

- Las columnas se resuelven por NOMBRE de cabecera, nunca por indice. La unica
  excepcion, obligada, son las columnas gemelas sin cabecera de la segunda mensula:
  se resuelven como la columna sin cabecera ADYACENTE POR EL LADO EXTERIOR (alejandose
  de la banda central) de su columna con cabecera.
- Nada se descarta en silencio: una cabecera desconocida va a NO_RECONOCIDO y las
  filas que no son un apoyo —hitos de estacion, notas al pie— salen con su motivo.
- El origen escribe en portugues y en espanol; a la base de datos llega INGLES. La
  tabla 'translations' es la unica que conoce las grafias del origen.
"""

from __future__ import annotations

import collections
import re
from decimal import Decimal, InvalidOperation

from openpyxl.utils import get_column_letter

from workbook_common import (
    MAX_COLS,
    MAX_ROWS_TRACK,
    code_tokens,
    is_blank,
    is_noise,
    norm_header,
    squash,
)

SYNOPTIC_FORMAT = "synoptic"

# La cabecera esta en las primeras filas, debajo del titulo del caderno y de la fila que
# anuncia 'VIA 2' / 'VIA 1'. Se localiza por contenido, no por numero de fila.
HEADER_SEARCH_ROWS = 12

# Un texto en la columna del apoyo que NO es un apoyo: 'Inic. Est. Campo Alegre (00+966)',
# 'Fin. Ponte (02+059)', 'Cruce Vial', 'Existente'. Son hitos del trazado y se distinguen
# porque no traen KP; el texto solo confirma que lo son.
LANDMARK_RE = re.compile(r"^(inic|fin|cruce|existente)", re.IGNORECASE)

# Campos multivalor del sinoptico y la entidad LOV a la que va cada campo. 'EQUIPMENT'
# no es un campo del perfil: es la columna 'Equipamentos associados', que mezcla anclajes
# con aparatos de seccionamiento y se reparte por codigo con la tabla 'routing'.
MULTI_FIELDS = ("SECTIONING", "EQUIPMENT")
SINGLE_LOV_FIELDS = ("SUPPORT_TYPE", "POLE_TYPE", "FOUNDATION", "ASSEMBLY_CONFIGURATION")
CANTILEVER_FIELDS = ("CANTILEVER_TYPE", "STAGGER", "CW_HEIGHT")

FIELD_ENTITY = {
    "SECTIONING": "Sectioning",
    "ANCHORAGE": "Anchorage",
    "SECTIONING_FEEDING": "DisconnectorFunction",
    "SUPPORT_TYPE": "SupportType",
    "POLE_TYPE": "PoleType",
    "FOUNDATION": "Foundation",
    "ASSEMBLY_CONFIGURATION": "AssemblyConfiguration",
    "CANTILEVER_TYPE": "CantileverType",
}

Block = collections.namedtuple(
    "Block", "marker first last outward single multi twins unmapped labels ignored")
SynopticLayout = collections.namedtuple(
    "SynopticLayout", "sheet rows header_index data_end blocks centre")


def is_synoptic(declared) -> bool:
    """True si el EP declara 'format: synoptic' en topology.yml."""
    return bool(declared) and squash(declared.get("format")).lower() == SYNOPTIC_FORMAT


def synoptic_config(cfg) -> dict:
    scfg = cfg.get("synoptic")
    if not scfg:
        raise KeyError("aliases.yml no tiene la seccion 'synoptic'")
    return scfg


def cell(row, index):
    return row[index] if index is not None and index < len(row) else None


# --------------------------------------------------------------------------------
# Plano de la hoja
# --------------------------------------------------------------------------------

def synoptic_layout(ws, ep, cfg, report):
    """Plano de la hoja: filas, cabecera, los dos bloques y la banda central.

    ``report`` es el acumulador del generador que llama (Master o Catalogue): los dos
    tienen ``unrecognised(kind, value, ep, sheet, row)``, que es lo unico que se usa.
    Devuelve None si la hoja no tiene cabecera o no tiene los dos bloques.
    """
    scfg = synoptic_config(cfg)
    sheet = ws.title
    rows = list(ws.iter_rows(min_row=1, max_row=min(ws.max_row, MAX_ROWS_TRACK),
                             max_col=MAX_COLS, values_only=True))

    marker = squash(scfg["header_marker"]).upper()
    header_index = next(
        (index for index, row in enumerate(rows[:HEADER_SEARCH_ROWS])
         if any(squash(value).upper() == marker for value in row)), None)
    if header_index is None:
        report.unrecognised("hoja sinoptico sin cabecera", sheet, ep, sheet, 1)
        return None
    header = rows[header_index]
    header_number = header_index + 1

    # Los marcadores 'VIA 2' / 'VIA 1' de la propia fila de cabecera parten la hoja: a la
    # izquierda del primero esta un bloque y a la derecha del segundo el otro. Lo de en
    # medio es la banda central, sin cabecera.
    wanted = [squash(m).upper() for m in scfg["block_markers"]]
    positions = {}
    for index, value in enumerate(header):
        text = squash(value).upper()
        if text in wanted:
            positions.setdefault(text, index)
    if len(positions) != 2:
        report.unrecognised("hoja sinoptico sin los dos bloques de via",
                            ", ".join(wanted), ep, sheet, header_number)
        return None
    (left_marker, left_pos), (right_marker, right_pos) = sorted(
        positions.items(), key=lambda item: item[1])

    blocks = {
        left_marker: resolve_block(header, left_marker, 0, left_pos - 1, -1,
                                   scfg, ep, sheet, header_number, report),
        right_marker: resolve_block(header, right_marker, right_pos + 1, len(header) - 1, +1,
                                    scfg, ep, sheet, header_number, report),
    }
    if any(block is None for block in blocks.values()):
        return None

    centre = list(range(left_pos + 1, right_pos))
    data_end = find_data_end(rows, header_index, blocks.values())
    return SynopticLayout(sheet, rows, header_index, data_end, blocks, centre)


def resolve_block(header, marker, first, last, outward, scfg, ep, sheet, header_number, report):
    """Cabecera -> indices de UN bloque, por nombre.

    ``outward`` es +1 para el bloque de la derecha y -1 para el de la izquierda: el
    sentido en el que la hoja se aleja de la banda central, que es donde estan las
    columnas gemelas de la segunda mensula.
    """
    columns = {norm_header(k): v for k, v in scfg["columns"].items()}
    unmapped_headers = {norm_header(v) for v in scfg.get("unmapped", [])}
    ignored = {norm_header(v) for v in scfg.get("ignored", [])}
    twin_fields = set(scfg.get("twins", []))

    seen: dict[str, list[int]] = collections.defaultdict(list)
    unmapped: dict[str, int] = {}
    skipped: set[int] = set()
    for index in range(first, last + 1):
        name = norm_header(header[index])
        if not name:
            continue
        if name in ignored:
            skipped.add(index)
            continue
        if name in columns:
            seen[columns[name]].append(index)
        elif name in unmapped_headers:
            unmapped.setdefault(squash(header[index]), index)
        else:
            report.unrecognised("cabecera sinoptico", name, ep, sheet, header_number)

    single, multi, twins = {}, {}, {}
    for field, indexes in seen.items():
        # Con la cabecera repetida ('Equipamentos associados' dos veces), la principal es
        # la mas cercana a la banda central y la otra es su gemela.
        inner = max(indexes) if outward < 0 else min(indexes)
        single[field] = inner
        # Sin cabecera repetida, la gemela es la columna SIN cabecera pegada por fuera.
        candidate = inner + outward
        twin = None
        if len(indexes) > 1:
            twin = min((index for index in indexes if index != inner),
                       key=lambda index: abs(index - inner))
        elif field in twin_fields and first <= candidate <= last \
                and not norm_header(header[candidate]):
            twin = candidate
        if field in MULTI_FIELDS:
            # Un campo multivalor lee TODAS sus columnas: la funcion del apoyo
            # ('Tipología') y el seccionamiento ('Seccionamento') van los dos a SECTIONING.
            multi[field] = sorted(set(indexes) | ({twin} if twin is not None else set()))
        elif field in twin_fields and twin is not None:
            twins[field] = twin

    for required in ("PROFILE_ID", "KP"):
        if required not in single:
            report.unrecognised(f"bloque sin columna {required}", marker, ep, sheet, header_number)
            return None
    labels = {index: squash(header[index]) for index in range(first, last + 1)}
    return Block(marker, first, last, outward, single, multi, twins, unmapped, labels, skipped)


def block_is_blank(row, block: Block) -> bool:
    """True si el bloque no trae nada en esta fila (la columna A, que siempre lleva
    un 1, esta en 'ignored' y no cuenta)."""
    return all(is_blank(cell(row, index)) for index in range(block.first, block.last + 1)
               if index not in block.ignored)


def find_data_end(rows, header_index, blocks) -> int:
    """Indice de la primera fila en blanco (en los dos bloques) despues del ultimo apoyo.

    Debajo de los datos el sinoptico lleva notas al pie, que no son apoyos aunque
    escriban un codigo. Lo que las separa es una fila en blanco; lo que venga despues
    sale en DESCARTADOS con su fila, para que no se pierda en silencio.
    """
    seen_profile = False
    for offset in range(header_index + 1, len(rows)):
        row = rows[offset]
        if all(block_is_blank(row, block) for block in blocks):
            if seen_profile:
                return offset
            continue
        seen_profile = seen_profile or any(
            not is_blank(cell(row, block.single["PROFILE_ID"])) for block in blocks)
    return len(rows)


# --------------------------------------------------------------------------------
# Filas de un bloque
# --------------------------------------------------------------------------------

def parse_kp(value):
    """KP de una celda del sinoptico, en metros, o None si no es un KP.

    Ademas del numero admite la kilometracion con signo mas —'1+118,8' son 1118,8 m—
    y la coma decimal. '??' y '-' son huecos.
    """
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return Decimal(repr(value)) if isinstance(value, float) else Decimal(value)
    text = squash(value).replace(" ", "")
    if not text:
        return None
    match = re.fullmatch(r"([-+]?\d+)\+(\d+(?:[.,]\d+)?)", text)
    if match:
        kilometres, metres = match.groups()
        return Decimal(kilometres) * 1000 + Decimal(metres.replace(",", "."))
    try:
        return Decimal(text.replace(",", "."))
    except InvalidOperation:
        return None


def synoptic_records(layout: SynopticLayout, marker, cfg):
    """Las filas de un bloque, agrupadas por apoyo.

    Devuelve una lista de dicts con ``kind`` en {'profile', 'landmark', 'footnote'}.
    Las filas intermedias —las que no traen apoyo— pertenecen al apoyo ANTERIOR, igual
    que el vano en las hojas HR Track: es el vano hasta el apoyo siguiente, y el
    'Seccionamento' o el 'Disp. Transic.' escritos entre dos apoyos son del primero.
    """
    block = layout.blocks[squash(marker).upper()]
    single, twins = block.single, block.twins
    records, current = [], None

    for offset in range(layout.header_index + 1, len(layout.rows)):
        row = layout.rows[offset]
        number = offset + 1                       # el numero que ensena Excel
        raw_id = cell(row, single["PROFILE_ID"])
        code = squash(raw_id)

        if offset >= layout.data_end:
            if not block_is_blank(row, block):
                records.append({"kind": "footnote", "number": number, "code": code,
                                "detail": code or first_text(row, block)})
            continue

        if is_blank(raw_id) or is_noise(code):
            if current is not None:
                attach(current, row, number, block, layout, cfg)
            continue

        kp_raw = cell(row, single["KP"])
        kp = parse_kp(kp_raw)
        kind = "landmark" if kp is None and LANDMARK_RE.match(code) else "profile"
        current = {
            "kind": kind, "number": number, "code": code,
            "kp_raw": kp_raw, "kp": kp,
            "span": cell(row, single.get("SPAN")),
            "single": {field: cell(row, single.get(field)) for field in SINGLE_LOV_FIELDS},
            "rail_pole_distance": cell(row, single.get("RAIL_POLE_DISTANCE")),
            "multi": {field: [] for field in MULTI_FIELDS},
            "cantilever": [
                {field: cell(row, single.get(field)) for field in CANTILEVER_FIELDS},
                {field: cell(row, twins.get(field)) for field in CANTILEVER_FIELDS},
            ],
            "unmapped": [],
        }
        for field in MULTI_FIELDS:
            for index in block.multi.get(field, []):
                value = cell(row, index)
                if not is_blank(value):
                    current["multi"][field].append((squash(value), number, block.labels[index]))
        collect_unmapped(current, row, number, block, layout)
        records.append(current)

    return records


def attach(current, row, number, block: Block, layout: SynopticLayout, cfg):
    """Lo que trae una fila intermedia se apunta al apoyo anterior."""
    single = block.single
    if is_blank(current.get("span")):
        current["span"] = cell(row, single.get("SPAN"))
    for field in MULTI_FIELDS:
        for index in block.multi.get(field, []):
            value = cell(row, index)
            if not is_blank(value):
                current["multi"][field].append((squash(value), number, block.labels[index]))
    collect_unmapped(current, row, number, block, layout)


def collect_unmapped(current, row, number, block: Block, layout: SynopticLayout):
    for label, index in block.unmapped.items():
        value = cell(row, index)
        if not is_blank(value):
            current["unmapped"].append((label, squash(value), number))
    # La banda central no tiene cabecera. Sus textos ('CAT D', 'CAT R') se conservan con
    # la letra de su columna; los marcadores numericos (un 1 o un 0 en cada fila) no
    # dicen nada y se ignoran.
    for index in layout.centre:
        value = cell(row, index)
        text = squash(value)
        if text and not is_noise(text):
            current["unmapped"].append((f"centro {get_column_letter(index + 1)}", text, number))


def first_text(row, block: Block) -> str:
    for index in range(block.first, block.last + 1):
        value = cell(row, index)
        if index not in block.ignored and not is_blank(value):
            return squash(value)
    return ""


# --------------------------------------------------------------------------------
# Traduccion al catalogo
# --------------------------------------------------------------------------------

def translate(field, text, cfg):
    """Codigo del catalogo para una grafia del origen, o None si no esta declarada.

    Devuelve '' cuando la tabla dice que esa grafia NO es un codigo ('S' es la
    suspension simple, sin seccionamiento), que no es lo mismo que no estar declarada.
    """
    table = (synoptic_config(cfg).get("translations") or {}).get(field) or {}
    upper = squash(text).upper()
    for raw, target in table.items():
        if squash(raw).upper() == upper:
            return "" if target is None else squash(target)
    return None


def is_hole(text: str) -> bool:
    """True si la celda no lleva codigo: vacia, '0', '-' o un error de formula.

    No es is_noise: en las hojas HR Track 'M1', 'M2' y 'M3' son subcabeceras de las
    columnas de mensula, pero en el sinoptico son maciços de verdad.
    """
    return is_blank(text) or squash(text).startswith("#")


def translate_single(field, raw, cfg) -> str:
    """Un campo de un solo codigo: hueco si esta vacio o es ruido, si no traducido."""
    text = squash(raw)
    if is_hole(text):
        return ""
    translated = translate(field, text, cfg)
    return text if translated is None else translated


Translated = collections.namedtuple("Translated", "codes rejected noise unrouted")


def translate_codes(field, cells, cfg) -> Translated:
    """Reparte las celdas de un campo multivalor en codigos del catalogo.

    ``cells`` son tuplas (valor, fila) o (valor, fila, cabecera de origen). Devuelve los
    codigos como pares (campo_destino, codigo) sin repetidos, mas lo que no ha entrado:
    ``rejected`` (erratas declaradas, como 'Anc'), ``noise`` (anotaciones como
    'Canton 396,4 m', que se conservan en NO_MAPEADO con su cabecera) y ``unrouted``
    (codigos de 'Equipamentos' que ninguna regla de 'routing' reparte).

    Primero se prueba la celda ENTERA en la tabla de traduccion —'Transição Catenaria'
    y 'Caix. Secc. e Con.' llevan espacios— y solo si no esta se parte con code_tokens,
    la misma particion que usan las hojas HR Track, y se traduce trozo a trozo.
    """
    scfg = synoptic_config(cfg)
    rejections = {squash(v).upper() for v in (scfg.get("rejections") or {}).get(field, [])}
    noise = (scfg.get("noise") or {}).get(field) or {}
    reassignment = (scfg.get("reassignment") or {}).get(field) or {}
    routing = scfg.get("routing") or {}

    codes, rejected, noisy, unrouted = [], [], [], []
    for raw, number, *label in cells:
        text = squash(raw)
        if is_hole(text):
            continue
        pattern = noise.get("regex")
        if pattern and re.fullmatch(pattern, text, re.IGNORECASE):
            noisy.append((text, number, label[0] if label else field))
            continue

        whole = translate(field, text, cfg)
        if whole is not None:
            pieces = [whole]
        else:
            pieces = []
            for token in code_tokens(text):
                translated = translate(field, token, cfg)
                pieces.append(token if translated is None else translated)

        for piece in pieces:
            piece = squash(piece)
            if not piece:
                continue
            if piece.upper() in rejections:
                rejected.append((piece, number))
                continue
            target, code = field, piece
            moved = lookup(reassignment, piece)
            if moved is not None:
                target, code = moved["to"], squash(moved["code"])
            elif field == "EQUIPMENT":
                target = next((destination for destination, accepted in routing.items()
                               if any(squash(v).upper() == piece.upper() for v in accepted)),
                              None)
                if target is None:
                    unrouted.append((piece, number))
                    continue
            if not any(t == target and c.upper() == code.upper() for t, c in codes):
                codes.append((target, code))

    return Translated(codes, rejected, noisy, unrouted)


def lookup(table: dict, code: str):
    upper = squash(code).upper()
    for raw, value in table.items():
        if squash(raw).upper() == upper:
            return value
    return None


def describe(entity, code, cfg) -> str:
    """Descripcion inglesa declarada para un codigo nuevo, o '' si no hay ninguna.

    El sinoptico no trae leyenda, asi que las descripciones se declaran en aliases.yml:
    una a una, o por patron para las familias numeradas ('C.F.21').
    """
    table = (synoptic_config(cfg).get("descriptions") or {}).get(entity) or {}
    for raw, value in table.items():
        if raw != "pattern" and squash(raw).upper() == squash(code).upper():
            return squash(value)
    for regex, template in (table.get("pattern") or {}).items():
        match = re.fullmatch(regex, squash(code))
        if match:
            return match.expand(template)
    return ""


def lov_columns(cfg) -> list[tuple[str, str]]:
    """(cabecera del origen, entidad LOV) para la hoja USO_TRACKS del catalogo."""
    scfg = synoptic_config(cfg)
    found = []
    for header, field in scfg["columns"].items():
        if field == "EQUIPMENT":
            found.append((f"sinoptico: {header}", "Anchorage / DisconnectorFunction"))
        elif field in FIELD_ENTITY:
            found.append((f"sinoptico: {header}", FIELD_ENTITY[field]))
    return found
